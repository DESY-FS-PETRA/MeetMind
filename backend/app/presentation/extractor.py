from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


_WHITESPACE_RE = re.compile(r"[ \t\r\f\v]+")
_GENERIC_SLIDE_LABEL_RE = re.compile(r"^(?:slide|page)\s*[-:#]?\s*\d+\s*$", re.IGNORECASE)
_GENERIC_TITLE_TEXTS = {
    "slide vocabulary cues",
    "presentation",
    "speaker notes",
}


@dataclass(frozen=True, slots=True)
class ExtractedSlide:
    slide_number: int
    title: str
    body_text: str
    notes_text: str = ""

    @property
    def combined_text(self) -> str:
        return "\n".join(
            part for part in (self.title, self.body_text, self.notes_text) if part.strip()
        ).strip()


@dataclass(frozen=True, slots=True)
class _TitleCandidate:
    text: str
    top: int
    left: int
    font_size: float
    order: int


class PresentationExtractionError(ValueError):
    """Raised when a presentation cannot be parsed safely."""


def _clean_text(value: str | None) -> str:
    if not value:
        return ""
    lines: list[str] = []
    for raw_line in str(value).replace("\x00", "").splitlines():
        line = _WHITESPACE_RE.sub(" ", raw_line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _first_line(text: str, fallback: str) -> str:
    for line in text.splitlines():
        clean = line.strip()
        if clean and not _GENERIC_SLIDE_LABEL_RE.fullmatch(clean):
            return clean[:300]
    return fallback


def _shape_texts(shape) -> Iterable[str]:
    """Yield visible text from text boxes, grouped shapes, and tables."""
    if getattr(shape, "has_text_frame", False):
        text = _clean_text(getattr(shape, "text", ""))
        if text:
            yield text

    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            for cell in row.cells:
                text = _clean_text(cell.text)
                if text:
                    yield text

    # GroupShape exposes a ``shapes`` collection. Other shape types generally
    # do not, so getattr keeps this compatible across python-pptx versions.
    child_shapes = getattr(shape, "shapes", None)
    if child_shapes is not None:
        for child in child_shapes:
            yield from _shape_texts(child)


def _shape_max_font_size(shape) -> float:
    sizes: list[float] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.font.size is not None:
                sizes.append(float(paragraph.font.size.pt))
            for run in paragraph.runs:
                if run.font.size is not None:
                    sizes.append(float(run.font.size.pt))
    return max(sizes, default=0.0)


def _is_title_candidate(text: str) -> bool:
    clean = _clean_text(text)
    if not clean or len(clean) > 220:
        return False
    lines = clean.splitlines()
    if len(lines) > 2:
        return False
    first = lines[0].strip()
    normalized = first.casefold().strip(" .:-")
    if _GENERIC_SLIDE_LABEL_RE.fullmatch(first):
        return False
    if normalized in _GENERIC_TITLE_TEXTS:
        return False
    if normalized.startswith("reference package for testing"):
        return False
    if normalized.startswith("these terms should bias"):
        return False
    return True


def _select_visual_title(slide, *, slide_height: int) -> str:
    """Select a title even when the deck does not use title placeholders.

    Many scientific decks are exported from custom templates and contain only
    ordinary text boxes. In those files ``slide.shapes.title`` is ``None``.
    We therefore prefer concise text boxes in the upper title band, ranked by
    font size and position, and only then fall back to other visible text.
    """
    candidates: list[_TitleCandidate] = []
    for order, shape in enumerate(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        text = _clean_text(getattr(shape, "text", ""))
        if not _is_title_candidate(text):
            continue
        candidates.append(
            _TitleCandidate(
                text=text.splitlines()[0].strip()[:300],
                top=int(getattr(shape, "top", 0) or 0),
                left=int(getattr(shape, "left", 0) or 0),
                font_size=_shape_max_font_size(shape),
                order=order,
            )
        )

    if not candidates:
        return ""

    title_band_limit = int(max(1, slide_height) * 0.23)
    title_band = [item for item in candidates if item.top <= title_band_limit]
    pool = title_band or candidates
    # Font size is the strongest cue inside the title band. A small top/left
    # preference breaks ties without allowing a footer or slide-number label
    # to win merely because it appears first in the XML.
    selected = max(
        pool,
        key=lambda item: (
            item.font_size,
            -item.top,
            -item.left,
            -item.order,
        ),
    )
    return selected.text


def extract_pptx(path: Path) -> list[ExtractedSlide]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise RuntimeError(
            "python-pptx is required for PowerPoint uploads. "
            "Run pip install -r requirements.txt"
        ) from exc

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise PresentationExtractionError(f"Could not read PowerPoint file: {exc}") from exc

    slides: list[ExtractedSlide] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        title_shape = getattr(slide.shapes, "title", None)
        if title_shape is not None:
            title = _clean_text(getattr(title_shape, "text", ""))
        if not title or _GENERIC_SLIDE_LABEL_RE.fullmatch(title):
            title = _select_visual_title(
                slide,
                slide_height=int(getattr(presentation, "slide_height", 0) or 0),
            )

        text_blocks: list[str] = []
        for shape in slide.shapes:
            for text in _shape_texts(shape):
                if title and text == title:
                    continue
                if _GENERIC_SLIDE_LABEL_RE.fullmatch(text):
                    continue
                if text not in text_blocks:
                    text_blocks.append(text)
        body = _clean_text("\n".join(text_blocks))

        notes = ""
        try:
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                notes = _clean_text(frame.text if frame is not None else "")
        except Exception:
            # Corrupt or unusual notes XML should not prevent slide ingestion.
            notes = ""

        slide_title = title or _first_line(body, f"Slide {number}")
        slides.append(
            ExtractedSlide(
                slide_number=number,
                title=slide_title,
                body_text=body,
                notes_text=notes,
            )
        )

    if not slides:
        raise PresentationExtractionError("The PowerPoint file contains no slides")
    return slides


def extract_pdf(path: Path) -> list[ExtractedSlide]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise RuntimeError(
            "pypdf is required for PDF uploads. Run pip install -r requirements.txt"
        ) from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise PresentationExtractionError(f"Could not read PDF file: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise PresentationExtractionError("Encrypted PDF files are not supported") from exc
        if not unlocked:
            raise PresentationExtractionError("Encrypted PDF files are not supported")

    slides: list[ExtractedSlide] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = _clean_text(page.extract_text() or "")
        except Exception as exc:
            raise PresentationExtractionError(
                f"Could not extract text from PDF page {number}: {exc}"
            ) from exc
        slides.append(
            ExtractedSlide(
                slide_number=number,
                title=_first_line(text, f"Slide {number}"),
                body_text=text,
                notes_text="",
            )
        )

    if not slides:
        raise PresentationExtractionError("The PDF contains no pages")
    return slides


def extract_presentation(path: Path, source_type: str) -> list[ExtractedSlide]:
    kind = source_type.casefold().lstrip(".")
    if kind == "pptx":
        return extract_pptx(path)
    if kind == "pdf":
        return extract_pdf(path)
    raise PresentationExtractionError("Only .pptx and .pdf files are supported")
