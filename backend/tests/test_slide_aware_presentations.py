from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from fastapi import UploadFile
from pptx import Presentation
from pptx.util import Inches

from app.presentation.service import PresentationService
from app.presentation.slide_context import SlideContextService
from app.storage.meeting_storage import MeetingStorage
from app.vocabulary.service import AdaptiveVocabularyService


class PresentationUploadTests(unittest.IsolatedAsyncioTestCase):
    async def test_pptx_text_notes_and_scientific_terms_are_persisted(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            database_path = root / "presentation.db"
            storage = MeetingStorage(
                database_url=f"sqlite:///{database_path.as_posix()}"
            )
            vocabulary = AdaptiveVocabularyService()
            service = PresentationService(
                storage=storage,
                vocabulary=vocabulary,
                root=root / "uploads",
            )
            try:
                deck = Presentation()
                first = deck.slides.add_slide(deck.slide_layouts[1])
                first.shapes.title.text = "Tomographic reconstruction"
                first.placeholders[1].text = "ASTRA Toolbox, SIRT and cone-beam geometry"
                first.notes_slide.notes_text_frame.text = (
                    "Compare filtered back projection with iterative reconstruction."
                )

                second = deck.slides.add_slide(deck.slide_layouts[1])
                second.shapes.title.text = "Near-field holography"
                second.placeholders[1].text = (
                    "Fresnel propagation, phase retrieval and coherence"
                )

                buffer = io.BytesIO()
                deck.save(buffer)
                buffer.seek(0)
                upload = UploadFile(filename="scientific_meeting.pptx", file=buffer)
                payload = await service.upload(upload)

                self.assertEqual(payload["slide_count"], 2)
                self.assertEqual(payload["slides"][0]["title"], "Tomographic reconstruction")
                self.assertIn("ASTRA Toolbox", payload["slides"][0]["keywords"])
                self.assertIn("SIRT", payload["slides"][0]["keywords"])
                self.assertIn(
                    "filtered back projection",
                    payload["slides"][0]["notes_text"],
                )
            finally:
                storage.database.dispose()


class SlideContextTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.storage = MeetingStorage(
            database_url=f"sqlite:///{(root / 'slides.db').as_posix()}"
        )
        self.storage.create_presentation(
            presentation_id="deck-1",
            filename="deck.pptx",
            source_type="pptx",
            title="Imaging meeting",
            storage_path=str(root / "deck.pptx"),
            slides=[
                {
                    "slide_number": 1,
                    "title": "Summer school introduction",
                    "body_text": "Location, schedule and accommodation",
                    "notes_text": "",
                    "keywords": [],
                },
                {
                    "slide_number": 2,
                    "title": "Near-field holography",
                    "body_text": (
                        "Fresnel propagation and phase retrieval for coherent "
                        "X-ray imaging"
                    ),
                    "notes_text": "",
                    "keywords": [
                        "near-field holography",
                        "Fresnel propagation",
                        "phase retrieval",
                    ],
                },
            ],
        )
        self.storage.create_meeting(
            "meeting-slides",
            "Slide test",
            "2026-07-23T00:00:00Z",
        )
        self.service = SlideContextService(self.storage)
        self.service.attach("meeting-slides", "deck-1")

    def tearDown(self) -> None:
        self.storage.database.dispose()
        self.temp_dir.cleanup()

    def test_raw_speech_moves_automatic_context_to_matching_slide(self) -> None:
        update = self.service.observe(
            "meeting-slides",
            "We now use Fresnel propagation for phase retrieval in near-field holography.",
        )
        self.assertIsNotNone(update)
        assert update is not None
        self.assertEqual(update.slide_number, 2)
        self.assertTrue(update.auto_match)
        self.assertIn("Fresnel propagation", update.slide_terms)

    def test_manual_slide_disables_automatic_matching(self) -> None:
        manual = self.service.activate_manual("meeting-slides", 1)
        self.assertFalse(manual.auto_match)
        update = self.service.observe(
            "meeting-slides",
            "Fresnel propagation phase retrieval near-field holography",
        )
        self.assertIsNone(update)

    def test_presentation_is_included_in_meeting_details(self) -> None:
        details = self.storage.get_meeting_details("meeting-slides")
        self.assertIsNotNone(details)
        assert details is not None
        self.assertEqual(details["presentation"]["slide_count"], 2)
        self.assertEqual(details["presentation"]["active_slide_number"], 1)


class SlideVocabularyTests(unittest.TestCase):
    def test_slide_terms_are_prioritized_and_prompt_remains_bounded(self) -> None:
        service = AdaptiveVocabularyService()
        meeting_id = "slide-vocab"
        service.start_meeting(meeting_id)
        update = service.set_slide_context(
            meeting_id,
            slide_number=4,
            slide_title="Tomographic reconstruction",
            slide_terms=[
                "ASTRA Toolbox",
                "SIRT",
                "cone-beam geometry",
                "filtered back projection",
            ],
            auto_match=True,
        )
        self.assertIsNotNone(update)
        assert update is not None
        self.assertEqual(update.current_slide_number, 4)
        self.assertEqual(update.active_terms[0], "ASTRA Toolbox")
        self.assertIn("SIRT", update.active_terms)
        self.assertLessEqual(len(update.active_terms), service.max_hotwords)
        context = service.context_for(meeting_id)
        self.assertLessEqual(
            len(context.hotwords.split()),
            service.max_hotword_words + len(context.active_terms),
        )


if __name__ == "__main__":
    unittest.main()
