from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from app.presentation.extractor import extract_pptx
from app.schemas.summary import MeetingSummaryContent, SlideSummary
from app.services.meeting_summary_service import MeetingSummaryService
from app.storage.meeting_storage import MeetingStorage
from app.transcription.result import TranscriptionResult, WordResult
from app.transcription.transcript_cleaner import TranscriptCleaner
from app.vocabulary.service import AdaptiveVocabularyService


class VisualTitleExtractionTests(unittest.TestCase):
    def test_custom_textbox_deck_uses_visual_title_not_slide_number(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "custom.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            number_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(1.0), Inches(0.3))
            number_box.text_frame.paragraphs[0].text = "Slide 2"
            number_box.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
            title_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.08), Inches(8.0), Inches(0.5))
            title_box.text_frame.paragraphs[0].text = "Motivation: synchrotron X-ray imaging"
            title_box.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
            body = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(7.0), Inches(4.0))
            body.text = "Near-field holography uses propagation-based phase contrast."
            deck.save(path)

            extracted = extract_pptx(path)

            self.assertEqual(extracted[0].title, "Motivation: synchrotron X-ray imaging")
            self.assertNotIn("Slide 2", extracted[0].body_text)


class ContextualCorrectionTests(unittest.TestCase):
    def _correct(self, title: str, terms: list[str], text: str) -> str:
        service = AdaptiveVocabularyService()
        service.start_meeting("precision")
        service.set_slide_context(
            "precision",
            slide_number=1,
            slide_title=title,
            slide_terms=terms,
            auto_match=True,
        )
        result = TranscriptionResult(
            text=text,
            words=tuple(WordResult(token, 0.2) for token in text.split()),
            average_confidence=0.2,
        )
        corrected, _ = service.correct_result("precision", result)
        cleaned = TranscriptCleaner().clean_words(list(corrected.words), final=True)
        return TranscriptCleaner.text(cleaned)

    def test_slide_specific_reference_errors_are_corrected(self) -> None:
        motivation = self._correct(
            "Motivation: synchrotron X-ray imaging",
            ["X-ray imaging", "phase retrieval"],
            "I introduce the modi-for synchrotron X-ray imaging.",
        )
        self.assertIn("the motivation for synchrotron X-ray imaging", motivation)

        ai = self._correct(
            "AI and machine learning for reconstruction",
            ["machine learning", "U-Net"],
            "I discuss artificial intelligence and l-learning for reconstruction.",
        )
        self.assertIn("machine learning", ai)
        self.assertNotIn("l-learning", ai)

    def test_distinct_qualified_facilities_and_error_examples_are_preserved(self) -> None:
        cleaner = TranscriptCleaner()
        text = (
            "PETRA III and PETRA IV beamlines. "
            "ASTRA Toolbox is not confused with ASTRA Toolbox."
        )
        words = [WordResult(token, 0.9) for token in text.split()]
        cleaned = TranscriptCleaner.text(cleaner.clean_words(words, final=True))
        self.assertIn("PETRA III and PETRA IV", cleaned)
        self.assertIn("not confused with ASTRA Toolbox", cleaned)

    def test_scientific_list_overlap_is_deduplicated(self) -> None:
        text = self._correct(
            "Ptychography and coherent diffraction imaging",
            ["PtyPy", "PyNX", "pyFAI", "cSAXS code"],
            "PtyPy PyNX PyPhy and the cSAX PtyPy, and the cSAXS code.",
        )
        self.assertEqual(text.count("PtyPy"), 1)
        self.assertEqual(text.count("cSAXS code"), 1)
        self.assertIn("pyFAI", text)


class SummaryPrecisionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.storage = MeetingStorage(
            database_url=f"sqlite:///{(root / 'summary.db').as_posix()}"
        )
        self.service = MeetingSummaryService(storage=self.storage)

    def tearDown(self) -> None:
        self.storage.database.dispose()
        self.temp_dir.cleanup()

    def test_incomplete_small_model_slide_summary_uses_full_spoken_section(self) -> None:
        sections = [
            {
                "slide_number": 4,
                "title": "Reconstruction pipeline",
                "transcript": (
                    "On slide 4, the raw frames need dark-field correction, flat-field "
                    "correction, cropping, normalization, Fresnel propagation, and "
                    "Paganin phase retrieval."
                ),
                "keywords": [
                    "dark-field correction",
                    "flat-field correction",
                    "Fresnel propagation",
                    "Paganin phase retrieval",
                ],
            }
        ]
        llm = [
            SlideSummary(
                slide_number=4,
                title="Slide 4",
                summary="The raw frames need dark-field correction.",
                key_terms=["dark-field correction"],
            )
        ]

        result = self.service._build_slide_summaries(llm, sections)

        self.assertEqual(result[0].title, "Reconstruction pipeline")
        self.assertIn("Fresnel propagation", result[0].summary)
        self.assertIn("Paganin phase retrieval", result[0].summary)

    def test_technical_terms_prefer_canonical_slide_terms_over_fragments(self) -> None:
        transcript = (
            "The sinogram uses ASTRA Toolbox, filtered back projection, and "
            "cone-beam geometry. ASTRA Toolbox is not confused with ASTA football."
        )
        sections = [
            {
                "slide_number": 5,
                "title": "Tomography reconstruction with ASTRA Toolbox",
                "transcript": transcript,
                "keywords": [
                    "ASTRA Toolbox",
                    "sinogram",
                    "filtered back projection",
                    "cone-beam geometry",
                ],
            }
        ]
        terms = self.service._build_technical_terms(
            ["Ray", "ASTA", "back projection", "cone-beam"],
            transcript,
            sections,
        )

        self.assertIn("ASTRA Toolbox", terms)
        self.assertIn("filtered back projection", terms)
        self.assertIn("cone-beam geometry", terms)
        self.assertNotIn("Ray", terms)
        self.assertNotIn("ASTA", terms)
        self.assertNotIn("back projection", terms)
        self.assertNotIn("cone-beam", terms)

    def test_canonical_low_confidence_terms_and_error_examples_are_not_uncertainties(self) -> None:
        transcript = (
            "The sinogram uses TomoPy and SIRT. ASTRA Toolbox is not confused "
            "with ASTA football."
        )
        words = [
            {"text": "sinogram", "probability": 0.14, "needs_review": True},
            {"text": "TomoPy", "probability": 0.29, "needs_review": True},
            {"text": "SIRT", "probability": 0.25, "needs_review": True},
            {"text": "ASTA", "probability": 0.19, "needs_review": True},
            {"text": "slide", "probability": 0.05, "needs_review": True},
        ]
        sections = [
            {
                "slide_number": 5,
                "title": "Tomography",
                "transcript": transcript,
                "keywords": ["sinogram", "TomoPy", "SIRT", "ASTRA Toolbox"],
            }
        ]
        result = self.service._build_uncertainties(
            [], transcript, [{"is_final": True, "words": words}],
            ["sinogram", "TomoPy", "SIRT", "ASTRA Toolbox"], sections,
        )
        self.assertEqual(result, [])


class ExplicitSlideAnchorTests(unittest.TestCase):
    def test_spoken_slide_numbers_override_noisy_auto_events(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            storage = MeetingStorage(
                database_url=f"sqlite:///{(root / 'alignment.db').as_posix()}"
            )
            try:
                storage.create_presentation(
                    presentation_id="deck",
                    filename="deck.pptx",
                    source_type="pptx",
                    title="Reference",
                    storage_path=str(root / "deck.pptx"),
                    slides=[
                        {"slide_number": 1, "title": "Title", "body_text": "", "notes_text": "", "keywords": []},
                        {"slide_number": 2, "title": "Motivation", "body_text": "", "notes_text": "", "keywords": ["X-ray imaging"]},
                        {"slide_number": 3, "title": "Setup", "body_text": "", "notes_text": "", "keywords": ["monochromator"]},
                    ],
                )
                storage.create_meeting("m", "Test", "2026-07-23T00:00:00Z")
                storage.attach_presentation("m", "deck")
                partial = "Intro. On slide 2, motivation for X-ray imaging begins."
                storage.upsert_transcript("m", partial, segment_id="rolling", revision=1, is_final=False)
                storage.set_active_slide("m", 3, auto_match=True, match_score=8.0, source="automatic")
                final = partial + " More motivation details. On slide 3, the monochromator is described."
                storage.upsert_transcript("m", final, segment_id="rolling", revision=2, is_final=True)

                sections = storage.build_slide_transcript_sections("m")

                by_slide = {item["slide_number"]: item["transcript"] for item in sections}
                self.assertIn("More motivation details", by_slide[2])
                self.assertNotIn("monochromator", by_slide[2])
                self.assertIn("monochromator", by_slide[3])
            finally:
                storage.database.dispose()


if __name__ == "__main__":
    unittest.main()
